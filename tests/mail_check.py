# Kikan - suivi des actifs et remise de materiel signee
# Copyright (C) 2026 Gabriel Chollet
# Developpe avec l'assistance d'une IA (Claude, Anthropic).
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

"""E-mails (B3 rappels, B4 rapport mensuel) : configuration, envois reels vers un faux serveur SMTP, cloisonnement
des destinataires, absence de doublons, pannes, planificateur.

Usage :
  docker exec -i -e TEST_ADMIN_USER=admin -e TEST_ADMIN_PASSWORD=... kikan-api python - < tests/mail_check.py

Le faux serveur SMTP tourne dans ce script (127.0.0.1:2525, meme conteneur que l'API). Cree des donnees ZZML et des
comptes sec_ml* (nettoyage : tests/cleanup_test_data.sql). Code de sortie 1 si un test echoue. ASCII uniquement.
"""
import email
import email.policy
import io
import json
import os
import socketserver
import sys
import threading
import urllib.error
import urllib.request
from datetime import date, datetime, timedelta

sys.path.insert(0, "/app")

B = os.environ.get("TEST_BASE_URL", "http://localhost:8000")
ADMIN = os.environ.get("TEST_ADMIN_USER", "admin")
ADMIN_PW = os.environ.get("TEST_ADMIN_PASSWORD", "")
PW = "Sec-Test-Passw0rd-2026"
PORT = 2525
fails = 0


def show(name, ok, detail=""):
    global fails
    if not ok:
        fails += 1
    print(("PASS " if ok else "FAIL ") + name + (" | " + str(detail)[:160] if detail else ""))


# ---------------------------------------------------------------- faux serveur SMTP
INBOX = []          # messages recus : email.message.EmailMessage
REFUSED = set()     # adresses refusees au RCPT


class Handler(socketserver.StreamRequestHandler):
    def handle(self):
        def w(text):
            self.wfile.write((text + "\r\n").encode())

        w("220 fake ESMTP")
        rcpts = []
        while True:
            line = self.rfile.readline()
            if not line:
                break
            cmd = line.decode(errors="replace").strip()
            up = cmd.upper()
            if up.startswith(("EHLO", "HELO")):
                w("250 fake")
            elif up.startswith("MAIL FROM"):
                rcpts = []
                w("250 ok")
            elif up.startswith("RCPT TO"):
                addr = cmd[cmd.index("<") + 1:cmd.index(">")].lower()
                if addr in REFUSED:
                    w("550 refused")
                else:
                    rcpts.append(addr)
                    w("250 ok")
            elif up == "DATA":
                w("354 go")
                raw = []
                while True:
                    row = self.rfile.readline()
                    if row in (b".\r\n", b""):
                        break
                    raw.append(row[1:] if row.startswith(b"..") else row)
                msg = email.message_from_bytes(b"".join(raw), policy=email.policy.default)
                msg["X-Test-Rcpt"] = ",".join(rcpts)
                INBOX.append(msg)
                w("250 queued")
            elif up in ("RSET", "NOOP"):
                w("250 ok")
            elif up == "QUIT":
                w("221 bye")
                break
            else:
                w("502 unsupported")


class Server(socketserver.ThreadingTCPServer):
    allow_reuse_address = True
    daemon_threads = True


def start_server():
    srv = Server(("127.0.0.1", PORT), Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return srv


def to(msg):
    return msg["X-Test-Rcpt"]


def body(msg):
    part = msg.get_body(preferencelist=("plain",))
    return part.get_content() if part else ""


def mails_for(addr):
    return [m for m in INBOX if to(m) == addr]


# ---------------------------------------------------------------- client HTTP
def call(method, path, payload=None, token=None):
    h = {"Content-Type": "application/json"}
    if token:
        h["Authorization"] = "Bearer " + token
    req = urllib.request.Request(B + path, data=json.dumps(payload).encode() if payload is not None else None,
                                 method=method, headers=h)
    try:
        r = urllib.request.urlopen(req, timeout=120)
        return r.status, r.read()
    except urllib.error.HTTPError as e:
        return e.code, e.read()


def login(user, pw):
    s, b = call("POST", "/auth/login", {"username": user, "password": pw})
    return json.loads(b)["access_token"] if s == 200 else None


ADM = login(ADMIN, ADMIN_PW)
if not ADM:
    sys.exit("Connexion admin impossible")


def j(method, path, payload=None, token=None):
    s, b = call(method, path, payload, token or ADM)
    try:
        return s, json.loads(b)
    except ValueError:
        return s, b


def account(name, role, mail, **scope):
    body_ = {"username": name, "password": PW, "role": role, "email": mail, **scope}
    s, r = j("POST", "/auth/accounts", body_)
    if s == 409:
        acc = next(a for a in j("GET", "/auth/accounts")[1] if a["username"] == name.lower())
        s, r = j("PUT", "/auth/accounts/%d" % acc["id"], {"actif": True, "role": role, "password": PW, "email": mail, **scope})
    return r


BASE = {"actif": True, "hote": "127.0.0.1", "port": PORT, "securite": "aucune", "utilisateur": "",
        "expediteur": "kikan@exemple.test", "verifier_cert": True, "liste": "liste@exemple.test", "heure": 7,
        "renouv": True, "renouv_jours": 30, "depart": True, "depart_jours": 14, "nonrestit": True, "nonrestit_sem": 2,
        "rapport": False, "rapport_jour": 1, "rapport_pj": "aucune"}


def cfg(**changes):
    return j("PUT", "/mail/config", {**BASE, **changes})


def run(**kw):
    return j("POST", "/mail/executer", {"rappels": kw.get("rappels", True), "rapport": kw.get("rapport", False)})


from sqlmodel import Session, delete, select  # noqa: E402
from app.db import engine  # noqa: E402
from app.models import ReminderSent, Setting  # noqa: E402

TODAY = date.today()


def reset_sent():
    with Session(engine) as s:
        for r in s.exec(select(ReminderSent)).all():
            s.delete(r)
        for key in ("mail_dernier_cycle", "mail_dernier_rapport"):
            row = s.get(Setting, key)
            if row:
                s.delete(row)
        s.commit()


def sent_count():
    with Session(engine) as s:
        return len(s.exec(select(ReminderSent)).all())


print("== Organisation et donnees de test ==")
ent = {}
for code in ("ZZML-A", "ZZML-B"):
    s, r = j("POST", "/entites", {"code": code, "nom": code})
    if s == 409:
        r = next(e for e in j("GET", "/entites")[1] if e["code"] == code)
    ent[code] = r["id"]
EA, EB = ent["ZZML-A"], ent["ZZML-B"]
account("sec_ml_it", "informatique", "it@exemple.test")
account("sec_ml_rha", "rh", "rha@exemple.test", portee_globale=False, entites=[EA])
account("sec_ml_rhb", "rh", "rhb@exemple.test", portee_globale=False, entites=[EB])
account("sec_ml_rh0", "rh", "rh0@exemple.test")
account("sec_ml_lec", "lecture", "lec@exemple.test", portee_globale=True)
IT, RA, LEC = login("sec_ml_it", PW), login("sec_ml_rha", PW), login("sec_ml_lec", PW)
show("comptes de test", all([IT, RA, LEC]))


def person(key, eid, statut, sortie):
    s, u = j("POST", "/users", {"matricule": "ZZML-" + key, "nom": "Nom" + key, "prenom": "ZZML" + key, "entite_id": eid,
                                "statut": "actif"})
    s2, a = j("POST", "/assets", {"type": "ZZML type", "marque": "ZZML", "modele": "M" + key, "numero_serie": "ZZML-SN-" + key,
                                  "identifiant_interne": "ZZML-ID-" + key, "entite_id": eid})
    j("POST", "/assignments", {"user_id": u["id"], "asset_ids": [a["id"]]})
    if sortie:  # le depart se declare par la page Sorties (statut en_depart) ; une sortie passee = materiel non restitue
        j("POST", "/lifecycle/depart/%d" % u["id"], {"date_sortie": sortie.isoformat()})
    return u


UA = person("A1", EA, "actif", TODAY + timedelta(days=5))
UB = person("B1", EB, "actif", TODAY + timedelta(days=3))
UX = person("A2", EA, "actif", TODAY - timedelta(days=20))
UF = person("A3", EA, "actif", TODAY + timedelta(days=200))   # depart lointain : pas de rappel
s, near = j("POST", "/assets", {"type": "ZZML type", "marque": "ZZML", "modele": "MR", "numero_serie": "ZZML-SN-R1",
                                "identifiant_interne": "ZZML-ID-R1", "entite_id": EA,
                                "echeance_renouvellement": (TODAY + timedelta(days=10)).isoformat()})
j("POST", "/assets", {"type": "ZZML type", "marque": "ZZML", "modele": "MR2", "numero_serie": "ZZML-SN-R2",
                      "identifiant_interne": "ZZML-ID-R2", "entite_id": EA,
                      "echeance_renouvellement": (TODAY + timedelta(days=400)).isoformat()})
show("donnees creees", all(isinstance(x, dict) and "id" in x for x in (UA, UB, UX, UF, near)))

print("== Configuration ==")
s, r = j("GET", "/mail/config")
show("par defaut : e-mails desactives", s == 200 and r["config"]["actif"] is False and r["pret"] is False, r)
show("le mot de passe SMTP n'est jamais expose", "password" not in json.dumps(r).lower().replace("mot_de_passe_defini", ""))
for name, payload, codes in [
    ("serveur invalide", {**BASE, "hote": "bad host!"}, (400, 422)),
    ("serveur avec saut de ligne", {**BASE, "hote": "a\r\nb"}, (400, 422)),
    ("expediteur invalide", {**BASE, "expediteur": "pas-une-adresse"}, (400, 422)),
    ("expediteur avec injection d'en-tete", {**BASE, "expediteur": "a@b.test\r\nBcc: x@y.test"}, (400, 422)),
    ("adresse de liste invalide", {**BASE, "liste": "x y@z"}, (400, 422)),
    ("port hors limites", {**BASE, "port": 0}, (422,)),
    ("securite inconnue", {**BASE, "securite": "tls13"}, (422,)),
    ("heure hors limites", {**BASE, "heure": 24}, (422,)),
    ("activation sans serveur", {**BASE, "hote": ""}, (400,)),
]:
    s, r = j("PUT", "/mail/config", payload)
    show("refuse : " + name, s in codes, s)
for who, tok in (("RH", RA), ("lecture", LEC)):
    for method, path in (("GET", "/mail/config"), ("PUT", "/mail/config"), ("POST", "/mail/test"), ("GET", "/mail/apercu"),
                         ("POST", "/mail/executer")):
        s, r = j(method, path, BASE if method == "PUT" else ({} if method == "POST" else None), token=tok)
        show("%s %s refuse a un compte %s" % (method, path, who), s == 403, s)

print("== Envois (faux serveur SMTP) ==")
srv = start_server()
s, r = cfg()
show("configuration enregistree", s == 200 and r["pret"] is True and r["config"]["hote"] == "127.0.0.1", r)
s, r = j("POST", "/mail/test", token=IT)
show("message de test envoye au compte connecte", s == 200 and r["destinataire"] == "it@exemple.test", r)
show("message de test recu", len(mails_for("it@exemple.test")) == 1 and "test" in mails_for("it@exemple.test")[0]["Subject"].lower())
INBOX.clear()
s, r = j("POST", "/mail/test")   # admin generique sans e-mail : repli sur l'adresse de liste
show("test sans e-mail de compte : adresse de liste", s == 200 and r["destinataire"] == "liste@exemple.test", r)
INBOX.clear()

reset_sent()
s, r = j("GET", "/mail/apercu")
dest = {d["destinataire"]: d["elements"] for d in r["rappels"]}
show("apercu : IT, RH A, RH B et liste, pas RH sans portee ni lecture",
     set(dest) == {"it@exemple.test", "rha@exemple.test", "rhb@exemple.test", "liste@exemple.test"}, sorted(dest))
show("apercu : rien n'a ete envoye", len(INBOX) == 0 and sent_count() == 0)

s, r = run()
show("execution : rappels envoyes", s == 200 and r["resultat"]["rappels"]["envoyes"] == 4, r.get("resultat"))
mi, ma, mb, ml = (mails_for(a) for a in ("it@exemple.test", "rha@exemple.test", "rhb@exemple.test", "liste@exemple.test"))
show("un seul recapitulatif par destinataire", all(len(x) == 1 for x in (mi, ma, mb, ml)), [len(x) for x in (mi, ma, mb, ml)])
bi, ba, bb, bl = body(mi[0]), body(ma[0]), body(mb[0]), body(ml[0])
show("informatique : renouvellement, departs et non restitue",
     "ZZML-ID-R1" in bi and "NomA1" in bi and "NomB1" in bi and "NomA2" in bi, bi[:200])
show("informatique : pas le depart lointain ni le renouvellement lointain",
     "NomA3" not in bi and "ZZML-ID-R2" not in bi)
show("RH A : ses salaries seulement", "NomA1" in ba and "NomA2" in ba and "NomB1" not in ba, ba[:200])
show("RH A : pas de renouvellement de materiel", "ZZML-ID-R1" not in ba and "renouvellement" not in ba.lower())
show("RH B : son depart seulement", "NomB1" in bb and "NomA1" not in bb and "NomA2" not in bb, bb[:200])
show("adresse de liste : tout comme l'informatique", "ZZML-ID-R1" in bl and "NomA1" in bl and "NomB1" in bl)
show("aucun message pour la RH sans portee ni la lecture", not mails_for("rh0@exemple.test") and not mails_for("lec@exemple.test"))
msg = mi[0]
show("en-tetes : expediteur, objet", msg["From"] == "kikan@exemple.test" and msg["Subject"].startswith("[Kikan] Rappels"), msg["Subject"])
show("pas de donnee sensible dans le message", "password" not in bi.lower() and "mot de passe" not in bi.lower())
n_sent = sent_count()
show("envois traces en base", n_sent > 0, n_sent)

print("== Pas de doublon, relances ==")
INBOX.clear()
s, r = run()
show("deuxieme execution : rien a renvoyer", s == 200 and r["resultat"]["rappels"]["envoyes"] == 0 and not INBOX, r.get("resultat"))
with Session(engine) as s_:
    for row in s_.exec(select(ReminderSent).where(ReminderSent.kind == "depart")).all():
        row.envoye_le = datetime.now() - timedelta(days=8)
        s_.add(row)
    s_.commit()
s, r = run()
bodies = " ".join(body(m) for m in INBOX)
show("relance des departs apres 7 jours, pas des autres", "NomA1" in bodies and "NomB1" in bodies and "ZZML-ID-R1" not in bodies
     and "NomA2" not in bodies, bodies[:200])

print("== Pannes ==")
reset_sent()
INBOX.clear()
REFUSED.add("rha@exemple.test")
s, r = run()
res = r["resultat"]["rappels"]
show("destinataire refuse : les autres recoivent", res["envoyes"] == 3 and len(res["erreurs"]) == 1 and not mails_for("rha@exemple.test"), res)
show("l'erreur ne contient pas le contenu du message", all("NomA" not in e for e in res["erreurs"]), res["erreurs"])
s, r = j("GET", "/mail/config")
show("derniere erreur memorisee", bool(r["derniere_erreur"]) and "rha@exemple.test" in r["derniere_erreur"], r["derniere_erreur"])
INBOX.clear()
REFUSED.clear()
s, r = run()
show("apres la panne : seul le destinataire manquant est servi", r["resultat"]["rappels"]["envoyes"] == 1 and len(mails_for("rha@exemple.test")) == 1,
     r.get("resultat"))
reset_sent()
INBOX.clear()
s, r = cfg(port=2599)   # personne n'ecoute
s, r = run()
res = r["resultat"]["rappels"]
show("serveur injoignable : echec propre, rien de marque comme envoye", s == 200 and res["envoyes"] == 0 and res["erreurs"] and sent_count() == 0, res)
s, r = j("POST", "/mail/test", token=IT)
show("test d'envoi vers un serveur injoignable : 502 lisible", s == 502 and "injoignable" in json.dumps(r).lower(), r)
cfg()
s, r = j("GET", "/health")
show("l'API reste disponible", s == 200)

print("== Desactivation ==")
s, r = cfg(actif=False)
s, r = run()
show("e-mails desactives : execution refusee", s == 409, r)
cfg()

print("== Rapport mensuel ==")
reset_sent()
INBOX.clear()
s, r = cfg(rapport=True, rapport_pj="xlsx", renouv=False, depart=False, nonrestit=False)
s, r = run(rappels=False, rapport=True)
show("rapport envoye a l'informatique et a la liste", s == 200 and r["resultat"]["rapport"]["envoyes"] == 2, r.get("resultat"))
mi = mails_for("it@exemple.test")
show("rapport : objet et contenu", len(mi) == 1 and "Rapport mensuel" in mi[0]["Subject"] and "Parc :" in body(mi[0]), mi and mi[0]["Subject"])
atts = [p for p in mi[0].iter_attachments()] if mi else []
show("rapport : piece jointe Excel valide", len(atts) == 1 and atts[0].get_filename().endswith(".xlsx"), [a.get_filename() for a in atts])
try:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(atts[0].get_content()))
    show("rapport : le classeur s'ouvre", wb.sheetnames == ["Detentions"], wb.sheetnames)
except Exception as exc:  # noqa: BLE001
    show("rapport : le classeur s'ouvre", False, exc)
show("rapport : pas de rappels dans ce message", "Rappels" not in mi[0]["Subject"])
show("pas de message pour la RH ni la lecture", not mails_for("rha@exemple.test") and not mails_for("lec@exemple.test"))
INBOX.clear()
cfg(rapport=True, rapport_pj="pdf", renouv=False, depart=False, nonrestit=False)
run(rappels=False, rapport=True)
atts = [p for m in mails_for("it@exemple.test") for p in m.iter_attachments()]
show("rapport : piece jointe PDF", len(atts) == 1 and atts[0].get_content().startswith(b"%PDF"), [a.get_filename() for a in atts])

print("== Planificateur ==")
from app import reminders  # noqa: E402

reset_sent()
INBOX.clear()
cfg(rapport=True, rapport_pj="aucune", rapport_jour=5, heure=7)
early = datetime(2026, 11, 10, 3, 0)
late = datetime(2026, 11, 10, 9, 0)
show("avant l'heure choisie : rien", reminders.tick(early) == {} or reminders.tick(early) is None and not INBOX)
show("avant l'heure choisie : aucun message", not INBOX)
res = reminders.tick(late)
show("apres l'heure : rappels et rapport partent", res and res["rappels"]["envoyes"] >= 1 and res["rapport"]["envoyes"] == 2, res)
n = len(INBOX)
res = reminders.tick(datetime(2026, 11, 10, 10, 0))
show("meme jour : rien de plus", not res and len(INBOX) == n, res)
res = reminders.tick(datetime(2026, 11, 11, 9, 0))
show("jour suivant : pas de rapport en double dans le mois", not res or "rapport" not in res, res)
res = reminders.tick(datetime(2026, 12, 6, 9, 0))
show("mois suivant (jour >= 5) : nouveau rapport", res and "rapport" in res and res["rapport"]["envoyes"] == 2, res)
reset_sent()
INBOX.clear()
res = reminders.tick(datetime(2026, 12, 3, 9, 0))
show("rapport pas avant le jour choisi", not res or "rapport" not in res, res)

print("== Donnees saisies hostiles ==")
reset_sent()
INBOX.clear()
cfg(rapport=False)
s, u = j("POST", "/users", {"matricule": "ZZML-INJ", "nom": "Dupont\r\nBcc: evil@exemple.test", "prenom": "ZZMLinj", "entite_id": EB,
                           "statut": "actif"})
s2, a = j("POST", "/assets", {"type": "ZZML type", "marque": "ZZML", "modele": "MI", "numero_serie": "ZZML-SN-INJ",
                              "identifiant_interne": "ZZML-ID-INJ", "entite_id": EB})
j("POST", "/assignments", {"user_id": u["id"], "asset_ids": [a["id"]]})
j("PUT", "/users/%d" % u["id"], {"matricule": "ZZML-INJ", "nom": "Dupont\r\nBcc: evil@exemple.test", "prenom": "ZZMLinj",
                                 "entite_id": EB, "statut": "actif", "date_sortie": (TODAY + timedelta(days=2)).isoformat()})
run()
mb = mails_for("rhb@exemple.test")
text = body(mb[0]) if mb else ""
show("un nom avec saut de ligne ne cree ni ligne ni en-tete", mb and "\nBcc:" not in text and mb[0]["Bcc"] is None
     and not mails_for("evil@exemple.test"), text[:240])

print("== Rappel des besoins de materiel (B7) ==")
reset_sent()
INBOX.clear()
cfg(renouv=False, depart=False, nonrestit=False, rapport=False, besoins=True, besoins_jours=14)
s, soon = j("POST", "/users", {"matricule": "ZZML-BS1", "nom": "NomBS1", "prenom": "ZZMLbs1", "entite_id": EA, "statut": "prevu",
                               "date_arrivee": (TODAY + timedelta(days=5)).isoformat()})
s, far = j("POST", "/users", {"matricule": "ZZML-BS2", "nom": "NomBS2", "prenom": "ZZMLbs2", "entite_id": EA, "statut": "prevu",
                              "date_arrivee": (TODAY + timedelta(days=60)).isoformat()})
j("POST", "/parametres/categories-besoin", {"nom": "ZZML categorie"})
s, b_soon = j("POST", "/besoins", {"user_id": soon["id"], "type": "ZZML categorie", "quantite": 2})
j("POST", "/besoins", {"user_id": far["id"], "type": "ZZML categorie"})
show("besoins crees pour les tests de rappel", s == 201, b_soon)
s, r = run()
bi = mails_for("it@exemple.test")
show("l'informatique recoit le rappel des besoins", len(bi) == 1 and "Besoins de materiel" in body(bi[0]) and "NomBS1" in body(bi[0])
     and "2 x ZZML categorie" in body(bi[0]), body(bi[0])[:200] if bi else r)
show("pas de rappel pour une arrivee lointaine", "NomBS2" not in body(bi[0]) if bi else False)
show("l'adresse de liste le recoit aussi", len(mails_for("liste@exemple.test")) == 1)
show("les RH ne recoivent pas ce rappel (traite par l'informatique)", not mails_for("rha@exemple.test") and not mails_for("rhb@exemple.test"))
INBOX.clear()
s, r = run()
show("pas de doublon le meme jour", not INBOX, r.get("resultat"))
s, r = j("GET", "/mail/apercu")
show("apercu coherent apres envoi", all("besoins" not in d["elements"] for d in r["rappels"]), r["rappels"])
j("POST", "/besoins/%d/statut" % b_soon["id"], {"statut": "annule"})
reset_sent()
INBOX.clear()
s, r = run()
show("un besoin annule n'est plus rappele", not any("NomBS1" in body(m) for m in INBOX))

print("== Audit ==")
s, rows = j("GET", "/audit?limit=200")
rows = rows if isinstance(rows, list) else rows.get("items", [])
actions = {r.get("action") for r in rows}
show("reglages, test et executions journalises", {"mail_reglages", "mail_test", "mail_execution_manuelle"} <= actions, sorted(a for a in actions if a and a.startswith("mail")))
show("le journal ne contient pas de contenu de message", all("NomA1" not in json.dumps(r) for r in rows))

srv.shutdown()
cfg(actif=False, renouv=False, depart=False, nonrestit=False, rapport=False)
print("")
print("ECHECS: %d" % fails)
sys.exit(1 if fails else 0)
