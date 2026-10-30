# Kikan - suivi des actifs et remise de materiel signee
# Copyright (C) 2026 Gabriel Chollet
# Developpe avec l'assistance d'une IA (Claude, Anthropic).
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

"""Lecture et ecriture de classeurs Excel (.xlsx), pour l'import, les modeles et l'export du rapport.

Un classeur est un fichier NON FIABLE (archive ZIP de XML) :
- archive verifiee avant toute lecture (nombre d'entrees, taille decompressee, macros, liens externes) ;
- XML lu avec defusedxml (bombes "billion laughs") ;
- lecture en mode lecture seule, VALEURS uniquement : les formules ne sont jamais evaluees ;
- a l'ecriture, toute cellule texte est ecrite comme du texte (un contenu "=..." ne s'execute jamais).
"""
import io
import zipfile
from datetime import date, datetime, time

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

if not openpyxl.DEFUSEDXML:  # sans defusedxml, openpyxl ne se protege pas des bombes XML
    raise RuntimeError("defusedxml est requis pour lire des classeurs Excel en securite")

MAX_ENTRIES = 2000
MAX_UNCOMPRESSED = 150 * 1024 * 1024
MAX_COLUMNS = 60
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class XlsxError(ValueError):
    """Classeur refuse (message affichable)."""


def is_xlsx(raw: bytes) -> bool:
    return raw[:4] == b"PK\x03\x04"


def is_legacy_office(raw: bytes) -> bool:
    return raw[:8] == OLE_MAGIC  # .xls ancien format ou classeur chiffre


def check_archive(raw: bytes) -> None:
    if is_legacy_office(raw):
        raise XlsxError("Format Excel ancien (.xls) ou classeur protege par mot de passe : enregistrez-le en .xlsx "
                        "sans mot de passe, ou en CSV.")
    try:
        infos = zipfile.ZipFile(io.BytesIO(raw)).infolist()
    except (zipfile.BadZipFile, ValueError, OSError) as exc:
        raise XlsxError("Classeur Excel illisible ou corrompu.") from exc
    if len(infos) > MAX_ENTRIES or sum(i.file_size for i in infos) > MAX_UNCOMPRESSED:
        raise XlsxError("Classeur refuse (archive anormalement volumineuse).")
    names = {i.filename for i in infos}
    if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
        raise XlsxError("Ce fichier n'est pas un classeur Excel (.xlsx).")
    lowered = [n.lower() for n in names]
    if any(n.endswith("vbaproject.bin") for n in lowered):
        raise XlsxError("Les classeurs avec macros sont refuses : enregistrez-les en .xlsx sans macro, ou en CSV.")
    if any(n.startswith("xl/externallinks/") for n in lowered):
        raise XlsxError("Le classeur contient des liens vers d'autres fichiers : collez les valeurs puis enregistrez-le.")


def _text(value) -> str:
    """Valeur de cellule -> texte, comme la lirait le CSV (dates ISO, entiers sans .0, booleens oui/non)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "oui" if value else "non"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0, 0) else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else format(value, "f").rstrip("0").rstrip(".")
    return str(value).strip()


def read_xlsx(raw: bytes, max_rows: int) -> tuple[list[dict], list[str], dict]:
    """(lignes, en-tetes, informations) comme read_csv. info : encodage, separateur, feuille, numeros de ligne."""
    check_archive(raw)
    try:
        values = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        formulas = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    except Exception as exc:  # noqa: BLE001  (openpyxl leve plusieurs types selon l'archive)
        raise XlsxError("Classeur Excel illisible ou corrompu.") from exc
    try:
        visible = [ws.title for ws in values.worksheets if ws.sheet_state == "visible"]
        if not visible:
            raise XlsxError("Le classeur ne contient aucune feuille visible.")
        sheet = visible[0]
        ws_v, ws_f = values[sheet], formulas[sheet]
        rows_v = ws_v.iter_rows(values_only=True, max_col=MAX_COLUMNS)
        rows_f = ws_f.iter_rows(values_only=True, max_col=MAX_COLUMNS)
        try:
            head_v, head_f = next(rows_v), next(rows_f)
        except StopIteration:
            raise XlsxError("La feuille est vide.") from None
        headers = [_text(h) for h in head_v]
        while headers and not headers[-1]:
            headers.pop()  # colonnes vides en bout de ligne
        if not any(headers):
            raise XlsxError("La premiere ligne de la feuille doit contenir les en-tetes de colonnes.")
        rows, numbers = [], []
        for n, (rv, rf) in enumerate(zip(rows_v, rows_f), start=2):
            if len(rows) > max_rows:
                break
            cells = [_text(c) for c in rv[:len(headers)]]
            cells += [""] * (len(headers) - len(cells))
            lost = [h for h, v, f in zip(headers, rv, rf) if v is None and isinstance(f, str) and f.startswith("=")]
            if not any(cells) and not lost:
                continue  # ligne vide
            row = {h: c for h, c in zip(headers, cells) if h}
            if lost:
                row["__formule"] = ", ".join(lost)
            rows.append(row)
            numbers.append(n)
        info = {"encodage": "Excel (.xlsx)", "separateur": "", "feuille": sheet,
                "feuilles_ignorees": len(values.sheetnames) - 1, "numeros_lignes": numbers}
        return rows, [h for h in headers if h], info
    finally:
        values.close()
        formulas.close()


# --- Ecriture ---

HEADER_FILL = PatternFill("solid", fgColor="2B4C7E")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def put_text(ws, row: int, col: int, value) -> None:
    """Cellule texte : ecrite comme du texte meme si elle commence par = + - @ (jamais interpretee)."""
    cell = ws.cell(row=row, column=col, value="" if value is None else str(value))
    cell.data_type = "s"
    cell.number_format = "@"


def _style_header(ws, headers: list[str]) -> None:
    for c, h in enumerate(headers, start=1):
        put_text(ws, 1, c, h)
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def _widths(ws, headers: list[str], rows: list[list], cap: int = 48) -> None:
    for c, h in enumerate(headers, start=1):
        longest = max([len(str(h))] + [len(str(r[c - 1])) for r in rows if c - 1 < len(r) and r[c - 1] is not None])
        ws.column_dimensions[get_column_letter(c)].width = min(max(longest + 3, 10), cap)


def _to_bytes(wb) -> bytes:
    wb.properties.creator = "Kikan"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_template(spec: dict, choices: dict[str, list[str]]) -> bytes:
    """Modele d'import : feuille Donnees (en-tetes, exemples, colonnes au format texte, listes) et feuille Aide."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Donnees"
    headers = [c["nom"] for c in spec["colonnes"]]
    _style_header(ws, headers)
    for r, example in enumerate(spec["exemples"], start=2):
        for c, v in enumerate(example, start=1):
            put_text(ws, r, c, v)
    for c in range(1, len(headers) + 1):
        for r in range(len(spec["exemples"]) + 2, 502):
            ws.cell(row=r, column=c).number_format = "@"  # texte : les zeros initiaux sont preserves
    for c, name in enumerate(headers, start=1):
        if name in choices:
            dv = DataValidation(type="list", formula1='"' + ",".join(choices[name]) + '"', allow_blank=True)
            dv.add(f"{get_column_letter(c)}2:{get_column_letter(c)}500")
            ws.add_data_validation(dv)
    _widths(ws, headers, spec["exemples"])
    aide = wb.create_sheet("Aide")
    for c, h in enumerate(["Colonne", "Obligatoire", "Contenu attendu"], start=1):
        put_text(aide, 1, c, h)
        aide.cell(row=1, column=c).fill, aide.cell(row=1, column=c).font = HEADER_FILL, HEADER_FONT
    for r, col in enumerate(spec["colonnes"], start=2):
        put_text(aide, r, 1, col["nom"])
        put_text(aide, r, 2, "oui" if col["obligatoire"] else "non")
        put_text(aide, r, 3, col["aide"])
    put_text(aide, len(spec["colonnes"]) + 3, 1, spec["description"])
    aide.column_dimensions["A"].width, aide.column_dimensions["B"].width, aide.column_dimensions["C"].width = 28, 12, 70
    return _to_bytes(wb)


def build_table(sheet_title: str, headers: list[str], rows: list[list], date_columns: set[int] = frozenset(),
                amount_columns: set[int] = frozenset()) -> bytes:
    """Tableau d'export : en-tetes, filtre automatique, volets figes, largeurs ; dates et montants (nombres calcules par
    Kikan) types, texte jamais interprete."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    _style_header(ws, headers)
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            if (c - 1) in date_columns and value:
                try:
                    cell = ws.cell(row=r, column=c, value=date.fromisoformat(str(value)[:10]))
                    cell.number_format = "yyyy-mm-dd"
                    continue
                except ValueError:
                    pass
            if (c - 1) in amount_columns and isinstance(value, (int, float)) and not isinstance(value, bool):
                ws.cell(row=r, column=c, value=value).number_format = "#,##0.00"
                continue
            put_text(ws, r, c, value)
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    _widths(ws, headers, rows)
    return _to_bytes(wb)
